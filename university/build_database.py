import json
import random
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent
EMAIL_DOMAIN = "uad.edu"
ECTS = 6

GRADE_SCALE = [
    ("A", 4.0, 93, 100), ("A-", 3.7, 90, 92), ("B+", 3.3, 87, 89), ("B", 3.0, 83, 86),
    ("B-", 2.7, 80, 82), ("C+", 2.3, 77, 79), ("C", 2.0, 73, 76), ("C-", 1.7, 70, 72),
    ("D+", 1.3, 67, 69), ("D", 1.0, 60, 66), ("F", 0.0, 0, 59),
]

TERMS = ["Fall 2024", "Spring 2025", "Fall 2025", "Spring 2026"]

PROGRAMS = {
    "Master of Global Management": ("Faculty of Business", [
        [("TAM 530", "Data Analytics and Digital Transformation in a Global World"),
         ("TGM 586", "Global Entrepreneurship and Sustainable Business"),
         ("TAM 557", "Global Marketing in a Digital Age"),
         ("TAM 542", "Global Leadership and Personal Development"),
         ("TAM 511", "Global Accounting: Managing with Numbers")],
        [("IAL 503", "Applied International Leadership"), ("FIN 560", "International Finance"),
         ("TGM 551", "Global Marketing Strategy"),
         ("TAM 505", "Navigating Global and Regional Business Environments"),
         ("TGM 507", "Global Organizational Consulting")],
        [("TGM 610", "Cross-Cultural Negotiation"), ("TGM 622", "Global Supply Chain Management"),
         ("TAM 615", "Strategic Management in Emerging Markets"),
         ("TGM 630", "Corporate Governance and Ethics"), ("TGM 690", "Capstone: Global Management Project")],
        [("TGM 640", "Global Risk Management"), ("TGM 650", "Digital Business Models"),
         ("TGM 660", "International Human Resource Management"), ("TGM 670", "Innovation Management"),
         ("TGM 699", "Master's Thesis")],
    ]),
    "Master of Business Administration": ("Faculty of Business", [
        [("MBA 501", "Managerial Economics"), ("MBA 510", "Financial Reporting and Analysis"),
         ("MBA 520", "Organizational Behavior"), ("MBA 530", "Marketing Management"),
         ("MBA 540", "Business Analytics")],
        [("MBA 550", "Corporate Finance"), ("MBA 560", "Operations Management"),
         ("MBA 570", "Strategic Human Resource Management"), ("MBA 580", "Business Ethics"),
         ("MBA 590", "Leadership and Change")],
        [("MBA 610", "Competitive Strategy"), ("MBA 620", "Entrepreneurship and Innovation"),
         ("MBA 630", "International Business"), ("MBA 640", "Negotiation"),
         ("MBA 690", "MBA Capstone Consulting Project")],
        [("MBA 650", "Mergers and Acquisitions"), ("MBA 660", "Digital Marketing Strategy"),
         ("MBA 670", "Project Management"), ("MBA 680", "Corporate Social Responsibility"),
         ("MBA 699", "MBA Final Project")],
    ]),
    "Bachelor of Software Engineering": ("Faculty of Computer Science", [
        [("CSE 101", "Introduction to Programming"), ("MTH 110", "Discrete Mathematics"),
         ("CSE 120", "Computer Systems Fundamentals"), ("ENG 101", "Academic English I"),
         ("MTH 120", "Calculus I")],
        [("CSE 150", "Object-Oriented Programming"), ("CSE 160", "Data Structures"),
         ("CSE 170", "Web Development Basics"), ("ENG 102", "Academic English II"),
         ("MTH 130", "Linear Algebra")],
        [("CSE 210", "Algorithms"), ("CSE 220", "Databases"),
         ("CSE 230", "Software Requirements Engineering"), ("CSE 240", "Operating Systems"),
         ("STA 201", "Probability and Statistics")],
        [("CSE 250", "Software Architecture"), ("CSE 260", "Computer Networks"),
         ("CSE 270", "Mobile Application Development"), ("CSE 280", "Software Testing"),
         ("CSE 290", "Team Software Project")],
    ]),
    "Master of Data Science": ("Faculty of Computer Science", [
        [("DSC 501", "Foundations of Data Science"), ("DSC 510", "Statistical Inference"),
         ("DSC 520", "Programming for Data Science (Python)"), ("DSC 530", "Data Management and SQL"),
         ("DSC 540", "Linear Algebra for Machine Learning")],
        [("DSC 550", "Machine Learning"), ("DSC 560", "Data Visualization"),
         ("DSC 570", "Big Data Platforms"), ("DSC 580", "Optimization Methods"),
         ("DSC 590", "Research Methods")],
        [("DSC 610", "Deep Learning"), ("DSC 620", "Natural Language Processing"),
         ("DSC 630", "Time Series Analysis"), ("DSC 640", "Data Ethics and Privacy"),
         ("DSC 690", "Capstone Project in Data Science")],
        [("DSC 650", "Reinforcement Learning"), ("DSC 660", "Computer Vision"),
         ("DSC 670", "Causal Inference"), ("DSC 680", "MLOps and Model Deployment"),
         ("DSC 699", "Master's Thesis")],
    ]),
    "Bachelor of Finance": ("Faculty of Economics & Finance", [
        [("ECO 101", "Principles of Microeconomics"), ("ACC 101", "Financial Accounting"),
         ("MTH 105", "Business Mathematics"), ("ENG 101", "Academic English I"),
         ("FIN 100", "Introduction to Finance")],
        [("ECO 102", "Principles of Macroeconomics"), ("ACC 102", "Managerial Accounting"),
         ("STA 110", "Business Statistics"), ("ENG 102", "Academic English II"),
         ("FIN 150", "Financial Markets and Institutions")],
        [("FIN 210", "Corporate Finance"), ("FIN 220", "Investments"),
         ("FIN 230", "Banking Operations"), ("ECO 210", "Econometrics I"),
         ("LAW 201", "Business Law")],
        [("FIN 240", "Financial Modelling"), ("FIN 250", "Risk Management"),
         ("FIN 260", "International Finance"), ("ECO 220", "Money and Banking"),
         ("ACC 210", "Intermediate Accounting")],
    ]),
    "Master of International Relations": ("Faculty of Social Sciences", [
        [("IRL 501", "Theories of International Relations"), ("IRL 510", "Global Political Economy"),
         ("IRL 520", "International Law"), ("IRL 530", "Research Design"),
         ("IRL 540", "Diplomacy and Statecraft")],
        [("IRL 550", "Security Studies"), ("IRL 560", "Central Asia in World Politics"),
         ("IRL 570", "International Organizations"), ("IRL 580", "Foreign Policy Analysis"),
         ("IRL 590", "Quantitative Methods")],
        [("IRL 610", "Conflict Resolution"), ("IRL 620", "Energy Geopolitics"),
         ("IRL 630", "Human Rights and Global Governance"), ("IRL 640", "Regional Integration"),
         ("IRL 690", "Master's Thesis Seminar")],
        [("IRL 650", "Public Diplomacy"), ("IRL 660", "Migration and Global Mobility"),
         ("IRL 670", "Environmental Politics"), ("IRL 680", "Negotiation Simulation"),
         ("IRL 699", "Master's Thesis")],
    ]),
    "Bachelor of Civil Engineering": ("Faculty of Engineering", [
        [("CIV 101", "Engineering Drawing"), ("MTH 120", "Calculus I"),
         ("PHY 101", "Physics I: Mechanics"), ("ENG 101", "Academic English I"),
         ("CHM 101", "General Chemistry")],
        [("CIV 150", "Statics"), ("MTH 140", "Calculus II"),
         ("PHY 102", "Physics II: Electricity and Magnetism"), ("ENG 102", "Academic English II"),
         ("CIV 160", "Surveying")],
        [("CIV 210", "Strength of Materials"), ("CIV 220", "Construction Materials"),
         ("CIV 230", "Fluid Mechanics"), ("MTH 210", "Differential Equations"),
         ("CIV 240", "Computer-Aided Design (CAD)")],
        [("CIV 250", "Structural Analysis"), ("CIV 260", "Soil Mechanics"),
         ("CIV 270", "Hydraulics"), ("CIV 280", "Transportation Engineering"),
         ("CIV 290", "Construction Management")],
    ]),
    "Bachelor of Graphic Design": ("Faculty of Arts & Design", [
        [("ART 101", "Drawing Fundamentals"), ("DES 101", "Design Principles"),
         ("ART 110", "Color Theory"), ("ENG 101", "Academic English I"),
         ("HIS 101", "History of Art")],
        [("DES 150", "Typography I"), ("DES 160", "Digital Imaging"),
         ("DES 170", "Visual Communication"), ("ENG 102", "Academic English II"),
         ("ART 150", "Photography Basics")],
        [("DES 210", "Typography II"), ("DES 220", "Brand Identity Design"),
         ("DES 230", "UI/UX Design"), ("DES 240", "Motion Graphics"),
         ("DES 250", "Print Production")],
        [("DES 260", "Packaging Design"), ("DES 270", "Illustration"),
         ("DES 280", "Interaction Design"), ("DES 290", "Portfolio Development"),
         ("ART 210", "Contemporary Art and Culture")],
    ]),
}

STUDENTS = [
    ("m240101", "Aziza Karimova", "Master of Global Management", "Fall 2024", "Active", 90),
    ("m240102", "Rustam Nazarov", "Master of Global Management", "Fall 2024", "Active", 81),
    ("m240103", "Shahzoda Ergasheva", "Master of Business Administration", "Fall 2024", "Active", 92),
    ("m240104", "Timur Abdullaev", "Bachelor of Software Engineering", "Fall 2024", "Active", 86),
    ("m240105", "Malika Usmonova", "Master of Data Science", "Fall 2024", "Active", 94),
    ("m240106", "Bobur Rakhmonov", "Bachelor of Finance", "Fall 2024", "On Leave", 74),
    ("m240107", "Kamila Sharipova", "Master of International Relations", "Fall 2024", "Active", 88),
    ("m240108", "Doniyor Kholmatov", "Bachelor of Civil Engineering", "Fall 2024", "Active", 79),
    ("m250201", "Alisher Juraev", "Bachelor of Software Engineering", "Spring 2025", "Active", 89),
    ("m250202", "Fariza Yunusova", "Master of Business Administration", "Spring 2025", "Active", 93),
    ("m250203", "Sardor Tursunov", "Master of Global Management", "Spring 2025", "Active", 84),
    ("m250204", "Madina Rakhimova", "Bachelor of Finance", "Spring 2025", "Active", 91),
    ("m250205", "Otabek Kalandarov", "Master of Data Science", "Spring 2025", "Active", 83),
    ("m250206", "Zarina Ismoilova", "Master of International Relations", "Spring 2025", "Active", 87),
    ("m250207", "Jamshidbek Mirzayev", "Bachelor of Civil Engineering", "Spring 2025", "Active", 72),
    ("m250208", "Nigina Khasanova", "Bachelor of Graphic Design", "Spring 2025", "Active", 95),
    ("m250209", "Bekzod Akhmedov", "Bachelor of Finance", "Spring 2025", "Active", 80),
    ("m250210", "Dilshoda Matyakubova", "Master of International Relations", "Spring 2025", "Active", 90),
    ("m250301", "Javohir Saidov", "Bachelor of Software Engineering", "Fall 2025", "Active", 78),
    ("m250302", "Laylo Mamatova", "Bachelor of Graphic Design", "Fall 2025", "Active", 88),
    ("m250303", "Sherzod Ibragimov", "Master of Global Management", "Fall 2025", "Active", 85),
    ("m250304", "Gulnora Tojiboeva", "Master of Business Administration", "Fall 2025", "Active", 89),
    ("m250305", "Ulugbek Normatov", "Master of Data Science", "Fall 2025", "Active", 76),
    ("m250306", "Sevara Qodirova", "Bachelor of Civil Engineering", "Fall 2025", "Active", 86),
]


def letter_for(score):
    for letter, _, lo, hi in GRADE_SCALE:
        if lo <= score <= hi:
            return letter
    return "F"


def build():
    rng = random.Random(2026)
    courses = {}
    for program, (_, sems) in PROGRAMS.items():
        for i, sem in enumerate(sems):
            for code, title in sem:
                courses.setdefault(code, {"code": code, "title": title, "ects": ECTS})
    students, grades = [], []
    for sid, name, program, enrolled, status, ability in STUDENTS:
        first, last = name.split(" ", 1)
        students.append({
            "student_id": sid, "full_name": name,
            "email": f"{first[0].lower()}{last.lower().replace(' ', '')}@{EMAIL_DOMAIN}",
            "faculty": PROGRAMS[program][0], "program": program,
            "level": "Master" if program.startswith("Master") else "Bachelor",
            "enrollment_date": enrolled, "status": status,
        })
        terms = TERMS[TERMS.index(enrolled):]
        if status == "On Leave":
            terms = terms[:2]
        for i, term in enumerate(terms):
            for code, _ in PROGRAMS[program][1][i]:
                score = max(45, min(100, round(rng.gauss(ability, 6))))
                grades.append({"student_id": sid, "term": term, "semester": f"Semester {i + 1}",
                               "course_code": code, "grade": letter_for(score), "score": score})
    return students, list(courses.values()), grades


HEAD_FILL = PatternFill("solid", fgColor="1F3A4D")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D2C3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_sheet(ws, headers, rows, widths):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


def write_excel(students, courses, grades, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    keys = ["student_id", "full_name", "email", "faculty", "program", "level", "enrollment_date", "status"]
    write_sheet(ws, ["Student ID", "Full Name", "University Email", "Faculty", "Program", "Level",
                     "Enrollment Date", "Status"],
                [[s[k] for k in keys] for s in students], [12, 24, 28, 30, 36, 10, 15, 11])

    wc = wb.create_sheet("Courses")
    write_sheet(wc, ["Course Code", "Course Title", "ECTS Credits"],
                [[c["code"], c["title"], c["ects"]] for c in courses], [13, 55, 12])

    wg = wb.create_sheet("Grades")
    rows = []
    for r, g in enumerate(grades, start=2):
        rows.append([
            g["student_id"], f'=IFERROR(VLOOKUP(A{r},Students!$A:$B,2,FALSE),"")', g["term"], g["semester"],
            g["course_code"], f'=IFERROR(VLOOKUP(E{r},Courses!$A:$C,2,FALSE),"")',
            f'=IFERROR(VLOOKUP(E{r},Courses!$A:$C,3,FALSE),0)', g["grade"], g["score"],
            f'=IFERROR(VLOOKUP(H{r},\'Grade Scale\'!$A:$B,2,FALSE),0)', f"=J{r}*G{r}",
            f'=IF(H{r}="F",0,G{r})',
        ])
    write_sheet(wg, ["Student ID", "Student Name", "Term", "Semester", "Course Code", "Course Title", "ECTS",
                     "Grade", "Final Score", "Grade Points", "Quality Points", "ECTS Earned"],
                rows, [12, 24, 13, 12, 12, 50, 7, 8, 11, 12, 13, 12])

    wsc = wb.create_sheet("Grade Scale")
    write_sheet(wsc, ["Letter", "Grade Points", "Min Score", "Max Score"],
                [list(x) for x in GRADE_SCALE], [10, 13, 11, 11])

    wsum = wb.create_sheet("GPA Summary")
    headers = ["Student ID", "Student Name", "Program"] + [f"{t} GPA" for t in TERMS] + \
              ["Cumulative GPA", "ECTS Attempted", "ECTS Earned"]
    rows = []
    for r, s in enumerate(students, start=2):
        row = [s["student_id"], f"=VLOOKUP(A{r},Students!$A:$B,2,FALSE)", f"=VLOOKUP(A{r},Students!$A:$E,5,FALSE)"]
        for t in TERMS:
            row.append(f'=IFERROR(ROUND(SUMIFS(Grades!$K:$K,Grades!$A:$A,$A{r},Grades!$C:$C,"{t}")'
                       f'/SUMIFS(Grades!$G:$G,Grades!$A:$A,$A{r},Grades!$C:$C,"{t}"),2),"")')
        row.append(f'=IFERROR(ROUND(SUMIFS(Grades!$K:$K,Grades!$A:$A,$A{r})/SUMIFS(Grades!$G:$G,Grades!$A:$A,$A{r}),2),"")')
        row.append(f"=SUMIFS(Grades!$G:$G,Grades!$A:$A,$A{r})")
        row.append(f"=SUMIFS(Grades!$L:$L,Grades!$A:$A,$A{r})")
        rows.append(row)
    write_sheet(wsum, headers, rows, [12, 24, 36, 15, 15, 15, 15, 15, 15, 13])
    wb.save(path)


def main():
    students, courses, grades = build()
    write_excel(students, courses, grades, OUT / "university_database.xlsx")
    data = {"terms": TERMS,
            "gradeScale": [{"letter": l, "points": p, "min": lo, "max": hi} for l, p, lo, hi in GRADE_SCALE],
            "students": students, "courses": courses, "grades": grades}
    html_path = OUT / "index.html"
    html = html_path.read_text(encoding="utf-8")
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    html = re.sub(r'(<script id="db" type="application/json">).*?(</script>)',
                  lambda m: m.group(1) + payload + m.group(2), html, flags=re.S)
    html_path.write_text(html, encoding="utf-8")
    print(f"{len(students)} students, {len(courses)} courses, {len(grades)} grade records")


if __name__ == "__main__":
    main()
